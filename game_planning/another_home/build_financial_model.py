"""Builds AnotherHome_Financial_Model.xlsx (Team 7, Intro to Game Planning 2026-2).

Run:  python3 build_financial_model.py   then recalculate with LibreOffice.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "AnotherHome_Financial_Model.xlsx")

F = "Arial"
BLUE = Font(name=F, color="0000FF")
BLACK = Font(name=F, color="000000")
GREEN = Font(name=F, color="008000")
BOLD = Font(name=F, bold=True)
BOLDW = Font(name=F, bold=True, color="FFFFFF")
TITLE = Font(name=F, bold=True, size=14)
ITAL = Font(name=F, italic=True, color="666666", size=9)
LINK = Font(name=F, color="0563C1", underline="single")
YELLOW = PatternFill("solid", fgColor="FFFF00")
HDR = PatternFill("solid", fgColor="5B3A29")
SUB = PatternFill("solid", fgColor="F3E3D3")
TOT = PatternFill("solid", fgColor="F2F2F2")
TOPB = Border(top=Side(style="thin", color="999999"))

USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
NUM = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'

wb = Workbook()


def style_row(ws, r, c1, c2, font=None, fill=None, border=None):
    for c in range(c1, c2 + 1):
        cell = ws.cell(r, c)
        if font: cell.font = font
        if fill: cell.fill = fill
        if border: cell.border = border


def header(ws, r, labels):
    for i, lab in enumerate(labels):
        c = ws.cell(r, 1 + i, lab)
        c.font = BOLDW; c.fill = HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def widths(ws, d):
    for col, w in d.items():
        ws.column_dimensions[col].width = w


def put(ws, r, c, v, fmt=None, font=None, fill=None):
    cell = ws.cell(r, c, v)
    if fmt: cell.number_format = fmt
    if font: cell.font = font
    if fill: cell.fill = fill
    return cell

# ------------------------------------------------------------------ Assumptions
A = wb.active
A.title = "Assumptions"
A["A1"] = "Another Home — Key Assumptions (edit blue cells only)"; A["A1"].font = TITLE
A["A2"] = "Blue = input · Yellow = key assumption to defend in the pitch · All money in USD"; A["A2"].font = ITAL
header(A, 4, ["General parameter", "Unit", "", "Value", "", "Note / source"])
general = [
    ("budget", "Production budget", "USD", 75000, USD, "Course sheet 'List of Groups', Team 7", True),
    ("dev_months", "Development duration", "months", 10, NUM, "See Roadmap", False),
    ("price", "Steam base price", "USD", 9.99, USD2, "Cozy sandbox builders: Townscaper ≈ $6, Tiny Glade ≈ $15", True),
    ("discount", "Average discount incl. sales & regional pricing", "%", 0.30, PCT, "Assumption: indie games sell much of their volume on sale", False),
    ("refund", "Refund rate", "%", 0.05, PCT, "Assumption: typical 3–8% for small indie titles", False),
    ("steam_fee", "Steam revenue share", "%", 0.30, PCT, "Steam standard 30% below $10M revenue", False),
    ("itch_fee", "itch.io revenue share", "%", 0.10, PCT, "itch.io default open revenue share 10%", False),
    ("dlc_price", "Furniture DLC pack price", "USD", 2.99, USD2, "Assumption: cosmetic pack (e.g. 'Autumn Cabin', 'Tiny Café')", False),
    ("ost_price", "Soundtrack price", "USD", 4.99, USD2, "Assumption: typical indie OST price", False),
    ("ost_attach", "Soundtrack attach rate", "%", 0.03, PCT, "Assumption", False),
    ("liveops", "Live-ops: monthly furniture drops, per year", "USD/yr", 12000, USD, "Keeps the original dev's 'new furniture every month' promise", False),
    ("mkt_year", "Ongoing marketing per year", "USD/yr", 4000, USD, "Creators, seasonal sales assets", False),
    ("hosting", "Tools, hosting & Steam fees per year", "USD/yr", 400, USD, "Steam networking for co-op is free; web hosting on itch.io", False),
    ("tax", "Income tax rate", "%", 0.10, PCT, "Simplified assumption — check local SME rate", False),
    ("disc", "Discount rate (NPV)", "%", 0.12, PCT, "Assumption for a small risky indie project", False),
]
N = {}
r = 5
for key, lab, unit, val, fmt, note, key_ass in general:
    put(A, r, 1, lab); put(A, r, 2, unit)
    put(A, r, 4, val, fmt, BLUE, YELLOW if key_ass else None)
    put(A, r, 6, note, font=ITAL)
    N[key] = f"Assumptions!$D${r}"
    r += 1
r += 1
header(A, r, ["Scenario parameter", "Unit", "Low", "Base", "High", "Note / source"])
r += 1
scen = [
    ("wish", "Steam wishlists at launch", "wishlists", (8000, 25000, 60000), NUM, "Browser version + Steam Next Fest + cozy creators; see Market sheet", True),
    ("conv", "Year-1 units per launch wishlist", "x", (0.40, 0.60, 0.90), '0.00"x"', "Rule of thumb: first-year sales ≈ 0.3–1.0× launch wishlists", True),
    ("y2", "Year-2 units as % of Year 1", "%", (0.35, 0.50, 0.60), PCT, "Long tail: seasonal sales, updates, DLC launches", False),
    ("y3", "Year-3 units as % of Year 2", "%", (0.50, 0.60, 0.70), PCT, "", False),
    ("dlc_attach", "DLC packs bought per player (avg)", "packs", (0.10, 0.20, 0.30), '0.00', "Cosmetic furniture packs; 2 packs released by Year 2", True),
    ("itch", "itch.io browser version: tips & supporter pack per year", "USD/yr", (1000, 2500, 5000), USD, "Browser game stays free (name-your-price) as a funnel", False),
]
S = {}
for key, lab, unit, vals, fmt, note, key_ass in scen:
    put(A, r, 1, lab); put(A, r, 2, unit)
    for i, v in enumerate(vals):
        put(A, r, 3 + i, v, fmt, BLUE, YELLOW if key_ass else None)
    put(A, r, 6, note, font=ITAL)
    S[key] = r
    r += 1
widths(A, {"A": 54, "B": 11, "C": 12, "D": 12, "E": 12, "F": 78})
A.freeze_panes = "A5"
SCN = [("Low", "C"), ("Base", "D"), ("High", "E")]


def sa(key, col):
    return f"Assumptions!${col}${S[key]}"

# ------------------------------------------------------------------ Development Cost
D = wb.create_sheet("Development Cost")
D["A1"] = "Development Cost — how the $75,000 budget is spent"; D["A1"].font = TITLE
D["A2"] = "Personnel cost = Monthly rate × Months × FTE share. Remote freelance rates (assumption). Producer & Director = the two students (unpaid)."; D["A2"].font = ITAL
header(D, 4, ["Category", "Item", "Monthly rate (USD)", "Months", "FTE share", "Cost (USD)", "% of budget", "Note"])
people = [
    ("Personnel", "Lead programmer (building tools, co-op, save system)", 2400, 9, 1.0, "Multi-room homes, resize/recolour, Steam co-op, Workshop"),
    ("Personnel", "Isometric 2D artist (furniture & rooms)", 2000, 8, 1.0, "≈ 400 new pieces + UI art; keeps the original soft pastel style"),
    ("Personnel", "Game designer (content & Cozy Requests)", 1800, 6, 1.0, "Room themes, optional requests, catalogue balance"),
    ("Personnel", "Sound designer / composer", 1500, 2, 1.0, "Satisfying placement sounds + ambient loops (players' favourite feature)"),
    ("Personnel", "QA & playtesting", 800, 3, 1.0, "Last 3 months incl. Steam Next Fest demo"),
]
r = 5; p1 = r
for cat, item, rate, m, fte, note in people:
    put(D, r, 1, cat); put(D, r, 2, item)
    put(D, r, 3, rate, USD, BLUE); put(D, r, 4, m, NUM, BLUE); put(D, r, 5, fte, '0.0', BLUE)
    put(D, r, 6, f"=C{r}*D{r}*E{r}", USD)
    put(D, r, 8, note, font=ITAL)
    r += 1
p2 = r - 1
put(D, r, 2, "Personnel subtotal", font=BOLD); put(D, r, 6, f"=SUM(F{p1}:F{p2})", USD)
style_row(D, r, 1, 8, fill=TOT); D.cell(r, 6).font = BOLD
pers = r; r += 2
other = [
    ("Tools & legal", "Software & asset licences", 600, "Engine/tool licences, fonts, SFX libraries"),
    ("Tools & legal", "Steam Direct fee", 100, "Official one-off fee per game (recoupable)"),
    ("Tools & legal", "Legal & company setup", 800, "Publishing agreement with original developer, EULA"),
    ("Tools & legal", "Localization (EN + 5 languages)", 2500, "ZH, JA, KO, RU, ES — cozy games sell strongly in Asia"),
    ("Marketing & launch", "Cozy creator / streamer campaign", 5000, "Many small cozy & decorating creators (TikTok, YouTube, Twitch)"),
    ("Marketing & launch", "Trailer & Steam page assets", 1500, "Announcement + launch trailer, capsule art, GIFs"),
    ("Marketing & launch", "Festivals & showcases", 1500, "Submissions: Wholesome Direct, Cozy Games showcases; Steam Next Fest is free"),
    ("Marketing & launch", "Paid social (TikTok / Reels)", 2000, "'Room makeover' timelapses"),
]
o1 = r
for cat, item, cost, note in other:
    put(D, r, 1, cat); put(D, r, 2, item); put(D, r, 6, cost, USD, BLUE); put(D, r, 8, note, font=ITAL)
    r += 1
o2 = r - 1
put(D, r, 2, "Tools, legal & marketing subtotal", font=BOLD); put(D, r, 6, f"=SUM(F{o1}:F{o2})", USD)
style_row(D, r, 1, 8, fill=TOT); D.cell(r, 6).font = BOLD
oth = r; r += 2
put(D, r, 2, "Contingency rate"); put(D, r, 5, 0.10, PCT, BLUE, YELLOW); put(D, r, 8, "Reserve for delays / scope changes (standard 10%)", font=ITAL)
cr = r; r += 1
put(D, r, 1, "Contingency"); put(D, r, 2, "Contingency reserve"); put(D, r, 6, f"=(F{pers}+F{oth})*E{cr}", USD)
cont = r; r += 2
put(D, r, 2, "TOTAL DEVELOPMENT & LAUNCH COST"); put(D, r, 6, f"=F{pers}+F{oth}+F{cont}", USD)
style_row(D, r, 1, 8, font=BOLD, fill=SUB, border=TOPB)
tot = r; DEV_TOTAL = f"'Development Cost'!$F${tot}"; r += 1
put(D, r, 2, "Budget (from Assumptions)"); put(D, r, 6, f"={N['budget']}", USD, GREEN)
bud = r; r += 1
put(D, r, 2, "Remaining budget (must be ≥ 0)", font=BOLD); put(D, r, 6, f"=F{bud}-F{tot}", USD)
put(D, r, 7, f'=IF(F{r}>=0,"OK — within budget","OVER BUDGET")', font=BOLD)
for rr in list(range(p1, p2 + 1)) + [pers] + list(range(o1, o2 + 1)) + [oth, cont, tot]:
    put(D, rr, 7, f"=IF($F${bud}=0,0,F{rr}/$F${bud})", PCT)
widths(D, {"A": 20, "B": 50, "C": 18, "D": 9, "E": 10, "F": 14, "G": 20, "H": 70})
D.freeze_panes = "A5"

# ------------------------------------------------------------------ Revenue
R = wb.create_sheet("Revenue")
R["A1"] = "Revenue Forecast — Price / Gross Revenue (3 years after Steam launch)"; R["A1"].font = TITLE
R["A2"] = "Premium Steam release ($9.99) + cosmetic furniture DLC + soundtrack + free browser version on itch.io (tips)."; R["A2"].font = ITAL
REV = {}
r = 4
for name, col in SCN:
    header(R, r, [f"{name.upper()} SCENARIO", "Unit", "Year 1", "Year 2", "Year 3", "3-yr total"])
    r += 1

    def line(key, label, unit, vals, fmt, total=True, bold=False, fill=None):
        global r
        put(R, r, 1, label); put(R, r, 2, unit)
        for i, f in enumerate(vals):
            put(R, r, 3 + i, f, fmt)
        if total:
            put(R, r, 6, f"=SUM(C{r}:E{r})", fmt)
        if bold: style_row(R, r, 1, 6, font=BOLD)
        if fill: style_row(R, r, 1, 6, fill=fill)
        REV[(name, key)] = r
        r += 1

    line("units", "Steam units sold", "copies",
         [f"={sa('wish', col)}*{sa('conv', col)}", f"=C{r}*{sa('y2', col)}", f"=D{r}*{sa('y3', col)}"], NUM)
    u = REV[(name, "units")]
    line("net_price", "Realized price per copy (after discounts & refunds)", "USD",
         [f"={N['price']}*(1-{N['discount']})*(1-{N['refund']})"] * 3, USD2, total=False)
    line("game", "Game sales, gross", "USD", [f"={L}{u}*{L}{u+1}" for L in "CDE"], USD)
    line("dlc", "Furniture DLC packs, gross", "USD", [f"={L}{u}*{sa('dlc_attach', col)}*{N['dlc_price']}" for L in "CDE"], USD)
    line("ost", "Soundtrack, gross", "USD", [f"={L}{u}*{N['ost_attach']}*{N['ost_price']}" for L in "CDE"], USD)
    line("itch", "itch.io browser version (tips, supporter pack)", "USD", [f"={sa('itch', col)}"] * 3, USD)
    g0 = REV[(name, "game")]
    line("gross", "GROSS REVENUE", "USD", [f"=SUM({L}{g0}:{L}{g0+3})" for L in "CDE"], USD, bold=True, fill=SUB)
    line("fee_steam", "– Steam revenue share", "USD", [f"=-({L}{g0}+{L}{g0+1}+{L}{g0+2})*{N['steam_fee']}" for L in "CDE"], USD)
    line("fee_itch", "– itch.io revenue share", "USD", [f"=-{L}{g0+3}*{N['itch_fee']}" for L in "CDE"], USD)
    g = REV[(name, "gross")]
    line("net", "NET REVENUE (to the team)", "USD", [f"={L}{g}+{L}{g+1}+{L}{g+2}" for L in "CDE"], USD, bold=True, fill=TOT)
    r += 1
widths(R, {"A": 50, "B": 10, "C": 14, "D": 14, "E": 14, "F": 14})

# ------------------------------------------------------------------ P&L
P = wb.create_sheet("P&L & Break-even")
P["A1"] = "Profit & Loss, Cash Flow and Break-even"; P["A1"].font = TITLE
P["A2"] = "Year 0 = development (budget spent). Years 1–3 = after Steam launch."; P["A2"].font = ITAL
PL = {}
r = 4
for name, col in SCN:
    header(P, r, [f"{name.upper()} SCENARIO", "Year 0 (dev)", "Year 1", "Year 2", "Year 3", "Total"])
    r += 1

    def pl(key, label, vals, bold=False, fill=None, fmt=USD, total=True, link=False):
        global r
        put(P, r, 1, label)
        for i, f in enumerate(vals):
            put(P, r, 2 + i, f, fmt, GREEN if (link and isinstance(f, str)) else None)
        if total: put(P, r, 6, f"=SUM(B{r}:E{r})", fmt)
        if bold: style_row(P, r, 1, 6, font=BOLD)
        if fill: style_row(P, r, 1, 6, fill=fill)
        PL[(name, key)] = r
        r += 1

    rn = REV[(name, "net")]
    pl("rev", "Net revenue", [0] + [f"=Revenue!{L}{rn}" for L in "CDE"], bold=True)
    pl("dev", "Development & launch cost", [f"=-{DEV_TOTAL}", 0, 0, 0], link=True)
    pl("live", "Live-ops: monthly furniture drops", [0] + [f"=-{N['liveops']}"] * 3, link=True)
    pl("mkt", "Ongoing marketing", [0] + [f"=-{N['mkt_year']}"] * 3, link=True)
    pl("host", "Tools, hosting & fees", [0] + [f"=-{N['hosting']}"] * 3, link=True)
    a = PL[(name, "dev")]; b = PL[(name, "host")]
    pl("costs", "TOTAL COSTS", [f"=SUM({L}{a}:{L}{b})" for L in "BCDE"], bold=True, fill=TOT)
    rv = PL[(name, "rev")]; cs = PL[(name, "costs")]
    pl("pbt", "Profit before tax", [f"={L}{rv}+{L}{cs}" for L in "BCDE"], bold=True)
    pb = PL[(name, "pbt")]
    pl("tax", "Income tax", [f"=-MAX(0,{L}{pb})*{N['tax']}" for L in "BCDE"])
    pl("np", "NET PROFIT / CASH FLOW", [f"={L}{pb}+{L}{pb+1}" for L in "BCDE"], bold=True, fill=SUB)
    npr = PL[(name, "np")]
    pl("cum", "Cumulative cash", [f"=B{npr}", f"=B{r}+C{npr}", f"=C{r}+D{npr}", f"=D{r}+E{npr}"], total=False)
    cum = PL[(name, "cum")]
    put(P, r, 1, "Break-even year", font=BOLD)
    put(P, r, 2, f'=IF(C{cum}>=0,"Year 1",IF(D{cum}>=0,"Year 2",IF(E{cum}>=0,"Year 3","After Year 3")))', font=BOLD)
    PL[(name, "be")] = r; r += 1
    put(P, r, 1, "ROI after 3 years (cumulative cash ÷ dev cost)", font=BOLD)
    put(P, r, 2, f"=IF({DEV_TOTAL}=0,0,E{cum}/{DEV_TOTAL})", PCT, BOLD)
    PL[(name, "roi")] = r; r += 1
    put(P, r, 1, "NPV of 3 years (at discount rate)", font=BOLD)
    put(P, r, 2, f"=B{npr}+NPV({N['disc']},C{npr}:E{npr})", USD, BOLD)
    PL[(name, "npv")] = r; r += 2
widths(P, {"A": 48, "B": 15, "C": 14, "D": 14, "E": 14, "F": 14})

# ------------------------------------------------------------------ Market
M = wb.create_sheet("Market Estimation")
M["A1"] = "Market Estimation — TAM / SAM / SOM"; M["A1"].font = TITLE
M["A2"] = "Approximate public figures — verify and add the source link before submitting."; M["A2"].font = ITAL
header(M, 4, ["Steam market (players)", "Value", "Note / source"])
mrows = [
    ("Steam monthly active users", 130_000_000, NUM, "Valve reported 132M MAU (2021); verify latest figure", True),
    ("Share who buy cozy / sandbox / building games", 0.05, PCT, "Assumption — 'Cozy' and 'Building' are large Steam tags", False),
    ("TAM — cozy & sandbox buyers", None, NUM, "Formula: MAU × share", False),
    ("Share interested in home / room decorating at < $15", 0.15, PCT, "Assumption", False),
    ("SAM — decorating-game buyers", None, NUM, "Formula: TAM × share", False),
    ("Share we can win in Year 1 with a $75k budget", 0.0154, PCT, "Assumption — calibrated to Base scenario", True),
    ("SOM — Year-1 copies", None, NUM, "Formula: SAM × share; compare with Revenue → Base units Y1", False),
]
r = 5
for lab, v, fmt, note, ka in mrows:
    put(M, r, 1, lab)
    if v is None:
        put(M, r, 2, f"=B{r-2}*B{r-1}", fmt, BOLD); M.cell(r, 1).font = BOLD
    else:
        put(M, r, 2, v, fmt, BLUE, YELLOW if ka else None)
    put(M, r, 3, note, font=ITAL)
    r += 1
put(M, r, 1, "Base-scenario Steam units in Year 1 (from Revenue)")
put(M, r, 2, f"=Revenue!C{REV[('Base', 'units')]}", NUM, GREEN)
r += 2
header(M, r, ["Demand signals (comparable games)", "Value", "Note / source"])
r += 1
sig = [
    ("Tiny Glade — copies sold in first month", 616000, "GameDiscoverCo newsletter, Oct 2024 — cozy builder, $15"),
    ("Townscaper — estimated copies sold", 610000, "Estimate 610k–1.1M, price under $6 — verify on Gamalytic/SteamDB"),
    ("diorama (Steam version by the original dev) — user reviews", 83, "95% positive (Very Positive) — proves demand for this exact game"),
    ("Another Home (itch.io) — player comments", 180, "≈ 180 comments on the itch.io page, monthly furniture updates"),
]
for lab, v, note in sig:
    put(M, r, 1, lab); put(M, r, 2, v, NUM, BLUE); put(M, r, 3, note, font=ITAL); r += 1
widths(M, {"A": 60, "B": 16, "C": 70})

# ------------------------------------------------------------------ Roadmap
G = WB = wb.create_sheet("Roadmap")
G["A1"] = "Roadmap — 10 months development + first year after launch"; G["A1"].font = TITLE
G["A2"] = "Change Start month / Duration (blue); the timeline fills automatically."; G["A2"].font = ITAL
MONTHS = 18
header(G, 4, ["Phase", "Key deliverable / milestone", "Start month", "Duration (months)"] + [f"M{i}" for i in range(1, MONTHS + 1)])
phases = [
    ("Pre-production", "Macro GDD, tech plan, deal with original developer", 1, 2),
    ("Core rebuild", "Multi-room homes, resize/recolour, save system", 2, 3),
    ("Content production", "≈ 400 new furniture pieces, 6 room themes, sounds", 4, 5),
    ("Steam page live", "Wishlist campaign starts (target 5k by M6)", 3, 1),
    ("Co-op & Cozy Requests", "Online co-op decorating, optional requests", 5, 3),
    ("Demo + Steam Next Fest", "Free demo, creator keys", 8, 1),
    ("Polish & localization", "5 languages, bug fixing, Steam Deck check", 8, 2),
    ("LAUNCH", "Steam release + browser version update on itch.io", 10, 1),
    ("Monthly furniture drops", "Free updates keep the community active", 11, 8),
    ("DLC pack #1 / #2", "Cosmetic furniture packs ($2.99)", 13, 5),
]
r = 5
for ph, dl, st, du in phases:
    put(G, r, 1, ph, font=BOLD if ph == "LAUNCH" else BLACK)
    put(G, r, 2, dl, font=ITAL)
    put(G, r, 3, st, font=BLUE); put(G, r, 4, du, font=BLUE)
    for m in range(1, MONTHS + 1):
        c = put(G, r, 4 + m, f'=IF(AND({m}>=$C{r},{m}<$C{r}+$D{r}),"■","")', font=Font(name=F, color="5B3A29", bold=True))
        c.alignment = Alignment(horizontal="center")
    r += 1
widths(G, {"A": 26, "B": 52, "C": 11, "D": 11})
for m in range(1, MONTHS + 1):
    G.column_dimensions[get_column_letter(4 + m)].width = 5

# ------------------------------------------------------------------ Comparables
C = wb.create_sheet("Comparable Games")
C["A1"] = "Comparable cozy decorating / sandbox games"; C["A1"].font = TITLE
C["A2"] = "Yellow cells: fill in from each store page before the pitch."; C["A2"].font = ITAL
header(C, 4, ["Game", "Developer", "Link", "Platform", "Price", "Reviews / rating (fill)", "What it does", "Lesson for Another Home"])
comps = [
    ("Another Home (our base game)", "lane (emiwa)", "https://emiwa.itch.io/home", "Browser (itch.io)", "Free",
     "Isometric room decorating sandbox, no goals, monthly furniture updates", "Proven cozy loop and loyal community — our starting point"),
    ("diorama", "lane (original dev)", "https://store.steampowered.com/app/2887460/diorama/", "Steam (PC)", "Check page",
     "Steam version: ≈ 1000 pieces, resize, weather & time, online & LAN co-op; 95% of 83 reviews positive", "Validates demand; our plan must add clear value on top"),
    ("Tiny Glade", "Pounce Light", "https://store.steampowered.com/app/2198150/Tiny_Glade/", "Steam (PC)", "$15",
     "Goal-less diorama castle builder; 616k copies in first month", "'No goals' sandbox can sell huge with a strong visual hook"),
    ("Townscaper", "Oskar Stålberg", "https://store.steampowered.com/app/1291340/Townscaper/", "PC, Switch, mobile", "≈ $6",
     "Toy-like town builder, no goals", "Cheap price + shareable screenshots = long tail"),
    ("Unpacking", "Witch Beam", "https://store.steampowered.com/app/1135690/Unpacking/", "PC, consoles", "≈ $20",
     "Zen puzzle of unpacking belongings into rooms", "Room + object story works; light goals help retention"),
]
r = 5
for row in comps:
    for i, v in enumerate(row[:5] + ("",) + row[5:]):
        c = put(C, r, 1 + i, v); c.alignment = Alignment(wrap_text=True, vertical="top")
    C.cell(r, 3).hyperlink = row[2]; C.cell(r, 3).font = LINK
    C.cell(r, 6).fill = YELLOW; C.cell(r, 6).font = BLUE
    r += 1
widths(C, {"A": 28, "B": 18, "C": 46, "D": 18, "E": 12, "F": 18, "G": 56, "H": 50})

# ------------------------------------------------------------------ Summary
Z = wb.create_sheet("Summary", 0)
Z["A1"] = "Another Home — Financial Model Summary (for Pitchdeck)"; Z["A1"].font = TITLE
Z["A2"] = "Team 7 · Intro to Game Planning 2026-2 · All values USD · Everything here is calculated — edit the Assumptions sheet"; Z["A2"].font = ITAL
header(Z, 4, ["Pricing (Price slide)", "Value", "", "", ""])
prices = [
    ("Another Home — Steam (PC / Steam Deck)", f"={N['price']}", USD2),
    ("Furniture DLC pack", f"={N['dlc_price']}", USD2),
    ("Original Soundtrack", f"={N['ost_price']}", USD2),
    ("Browser version on itch.io", "Free / name your price", None),
    ("Production budget", f"={N['budget']}", USD),
    ("Planned development + launch cost", f"={DEV_TOTAL}", USD),
]
r = 5
for lab, f, fmt in prices:
    put(Z, r, 1, lab); put(Z, r, 2, f, fmt, GREEN if fmt else None); r += 1
r += 1
header(Z, r, ["Scenario comparison (Gross Revenue & Targets slides)", "Low", "Base", "High", "Unit"])
r += 1
metrics = [
    ("Wishlists at launch", lambda n, c: f"={sa('wish', c)}", NUM, "wishlists"),
    ("Steam copies — Year 1", lambda n, c: f"=Revenue!C{REV[(n,'units')]}", NUM, "copies"),
    ("Steam copies — 3 years", lambda n, c: f"=Revenue!F{REV[(n,'units')]}", NUM, "copies"),
    ("Gross revenue — Year 1", lambda n, c: f"=Revenue!C{REV[(n,'gross')]}", USD, "USD"),
    ("Gross revenue — Year 2", lambda n, c: f"=Revenue!D{REV[(n,'gross')]}", USD, "USD"),
    ("Gross revenue — Year 3", lambda n, c: f"=Revenue!E{REV[(n,'gross')]}", USD, "USD"),
    ("Gross revenue — 3 years total", lambda n, c: f"=Revenue!F{REV[(n,'gross')]}", USD, "USD"),
    ("Net revenue — 3 years total", lambda n, c: f"=Revenue!F{REV[(n,'net')]}", USD, "USD"),
    ("Net profit incl. development — 3 years", lambda n, c: f"='P&L & Break-even'!F{PL[(n,'np')]}", USD, "USD"),
    ("ROI after 3 years", lambda n, c: f"='P&L & Break-even'!B{PL[(n,'roi')]}", PCT, "%"),
    ("NPV of 3 years", lambda n, c: f"='P&L & Break-even'!B{PL[(n,'npv')]}", USD, "USD"),
    ("Break-even", lambda n, c: f"='P&L & Break-even'!B{PL[(n,'be')]}", "General", ""),
]
SUMROW = {}
for lab, fn, fmt, unit in metrics:
    put(Z, r, 1, lab)
    for i, (n, c) in enumerate(SCN):
        cell = put(Z, r, 2 + i, fn(n, c), fmt, GREEN); cell.alignment = Alignment(horizontal="right")
    put(Z, r, 5, unit, font=ITAL)
    if "3 years total" in lab or "Break-even" in lab: style_row(Z, r, 1, 5, fill=SUB)
    SUMROW[lab] = r
    r += 1
r += 1
header(Z, r, ["Revenue mix — Base, 3 years", "USD", "% of gross", "", ""])
r += 1
m1 = r
for lab, key in (("Game sales (Steam)", "game"), ("Furniture DLC", "dlc"), ("Soundtrack", "ost"), ("itch.io browser tips", "itch")):
    put(Z, r, 1, lab); put(Z, r, 2, f"=Revenue!F{REV[('Base', key)]}", USD, GREEN); r += 1
put(Z, r, 1, "Total gross", font=BOLD); put(Z, r, 2, f"=SUM(B{m1}:B{r-1})", USD, BOLD)
for rr in range(m1, r + 1):
    put(Z, rr, 3, f"=IF($B${r}=0,0,B{rr}/$B${r})", PCT)
r += 2
header(Z, r, ["Budget split (Development Cost slide)", "USD", "% of budget", "", ""])
r += 1
for lab, rr in (("Team (personnel)", pers), ("Tools, legal & marketing", oth), ("Contingency", cont)):
    put(Z, r, 1, lab)
    put(Z, r, 2, f"='Development Cost'!F{rr}", USD, GREEN)
    put(Z, r, 3, f"='Development Cost'!G{rr}", PCT, GREEN)
    r += 1
widths(Z, {"A": 58, "B": 18, "C": 16, "D": 16, "E": 10})

# ------------------------------------------------------------------ How to use
L = wb.create_sheet("How to use")
for i, (t, f) in enumerate([
    ("How to use this model", TITLE),
    ("1. Change inputs only on 'Assumptions' and 'Development Cost' (blue text). Everything else recalculates.", BLACK),
    ("2. Yellow cells are key assumptions — be ready to explain them in the pitch.", BLACK),
    ("3. 'Summary' holds the numbers for the Pitchdeck: Price / Gross Revenue, Market Estimation, Targets, Development Cost.", BLACK),
    ("4. 'Market Estimation' and 'Comparable Games' use approximate public figures — verify and add sources.", BLACK),
    ("", BLACK),
    ("Colour code: blue = input · black = formula · green = link to another sheet · yellow fill = key assumption", ITAL),
    ("Example input: Assumptions → 'Steam wishlists at launch' Low 8,000 / Base 25,000 / High 60,000.", ITAL),
], start=1):
    put(L, i, 1, t, font=f)
L.column_dimensions["A"].width = 120

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows():
        for c in row:
            if c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)

wb.save(OUT)
print("saved", OUT)
